"""docx 解析：正文段落 + 表格逐行。"""
from pathlib import Path

from docx import Document

from services.parsing.base import ParsedBlock


class DocxParser:
    def parse(self, file_path: Path) -> list[ParsedBlock]:
        doc = Document(str(file_path))
        blocks: list[ParsedBlock] = []
        idx = 0
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                blocks.append(ParsedBlock(text=text, meta={"part": "body", "index": idx}))
                idx += 1
        for t, table in enumerate(doc.tables):
            for r, row in enumerate(table.rows):
                cells = [c.text.strip() for c in row.cells]
                text = " | ".join(c for c in cells if c)
                if text:
                    blocks.append(
                        ParsedBlock(text=text, meta={"part": f"table{t + 1}", "row": r + 1})
                    )
        return blocks
