"""Excel 解析：每个 sheet 逐行转成「表头: 值」文本，天然适合按行检索。"""
from pathlib import Path

from openpyxl import load_workbook

from riddleme.services.parsing.base import ParsedBlock


class ExcelParser:
    def parse(self, file_path: Path) -> list[ParsedBlock]:
        wb = load_workbook(str(file_path), read_only=True, data_only=True)
        blocks: list[ParsedBlock] = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(c).strip() if c is not None else "" for c in rows[0]]
            has_header = any(header)
            for r, row in enumerate(rows[1:] if has_header else rows, start=2 if has_header else 1):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if not cells:
                    continue
                if has_header:
                    pairs = [
                        f"{header[i]}: {str(c).strip()}"
                        for i, c in enumerate(row)
                        if c is not None and str(c).strip() and i < len(header) and header[i]
                    ]
                    text = " | ".join(pairs) if pairs else " | ".join(cells)
                else:
                    text = " | ".join(cells)
                blocks.append(
                    ParsedBlock(text=text, meta={"sheet": ws.title, "row": r, "source": file_path.name})
                )
        wb.close()
        return blocks
